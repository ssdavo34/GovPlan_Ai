"""
DOCX 문서 생성 서비스
"""
from typing import Dict, Optional
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


class DOCXGenerator:
    """DOCX 사업계획서 생성기"""

    def __init__(self):
        """초기화"""
        pass

    def generate(
        self,
        proposal_data: Dict,
        output_path: str,
        include_cover: bool = True
    ) -> str:
        """
        사업계획서 DOCX 생성

        Args:
            proposal_data: 사업계획서 데이터
            output_path: 출력 파일 경로
            include_cover: 표지 포함 여부

        Returns:
            생성된 DOCX 파일 경로
        """

        # 새 문서 생성
        doc = Document()

        # 한글 폰트 설정 (한글 문서용)
        self._set_document_font(doc, font_name='맑은 고딕')

        # 표지 추가
        if include_cover:
            self._add_cover_page(doc, proposal_data)
            doc.add_page_break()

        # 기업 정보 및 공고 정보
        self._add_info_section(doc, proposal_data)

        # 섹션별 내용 추가
        sections = proposal_data.get('sections', {})
        for section_name, section_content in sections.items():
            # 섹션 제목
            heading = doc.add_heading(section_name, level=1)
            heading.paragraph_format.space_before = Pt(12)
            heading.paragraph_format.space_after = Pt(6)

            # 섹션 내용 추가
            self._add_section_content(doc, section_content)

            # 섹션 간 여백
            doc.add_paragraph()

        # 문서 저장
        doc.save(output_path)

        return output_path

    def _set_document_font(self, doc: Document, font_name: str = '맑은 고딕'):
        """문서 기본 폰트 설정"""
        try:
            # 기본 스타일의 폰트 설정
            style = doc.styles['Normal']
            font = style.font
            font.name = font_name
            font.size = Pt(10)

            # 한글 폰트 설정 (한글 문서용)
            element = style.element
            element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        except Exception as e:
            print(f"Warning: 폰트 설정 실패: {e}")

    def _add_cover_page(self, doc: Document, proposal_data: Dict):
        """표지 페이지 추가"""

        # 상단 여백
        for _ in range(5):
            doc.add_paragraph()

        # 제목
        title = proposal_data.get('title', '사업계획서')
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.runs[0]
        title_run.font.size = Pt(24)
        title_run.font.bold = True

        # 부제목 (지원사업명)
        if proposal_data.get('project_name'):
            subtitle_para = doc.add_paragraph(proposal_data['project_name'])
            subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle_run = subtitle_para.runs[0]
            subtitle_run.font.size = Pt(14)

        # 여백
        for _ in range(3):
            doc.add_paragraph()

        # 기업 정보 테이블
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'

        # 테이블 데이터
        table_data = [
            ('기업명', proposal_data.get('company_name', 'N/A')),
            ('대표자', proposal_data.get('ceo_name', 'N/A')),
            ('사업자등록번호', proposal_data.get('business_number', 'N/A')),
            ('작성일', datetime.now().strftime('%Y년 %m월 %d일'))
        ]

        for i, (label, value) in enumerate(table_data):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = value

            # 셀 폰트 설정
            for cell in table.rows[i].cells:
                cell_para = cell.paragraphs[0]
                cell_run = cell_para.runs[0] if cell_para.runs else cell_para.add_run()
                cell_run.font.size = Pt(11)

    def _add_info_section(self, doc: Document, proposal_data: Dict):
        """기업 정보 및 공고 정보 섹션 추가"""

        # 기업 정보
        doc.add_heading('기업 정보', level=1)

        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'

        company_data = [
            ('기업명', proposal_data.get('company_name', 'N/A')),
            ('업종', proposal_data.get('industry_name', 'N/A')),
            ('소재지', proposal_data.get('region', 'N/A')),
            ('직원 수', f"{proposal_data.get('employee_count', 'N/A')}명"),
            ('기술 분야', ', '.join(proposal_data.get('technology_fields', [])) if proposal_data.get('technology_fields') else 'N/A')
        ]

        for i, (label, value) in enumerate(company_data):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = str(value)

        doc.add_paragraph()

        # 지원사업 정보
        doc.add_heading('지원사업 정보', level=1)

        table = doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'

        project_data = [
            ('사업명', proposal_data.get('project_name', 'N/A')),
            ('시행기관', proposal_data.get('agency', 'N/A')),
            ('지원 분야', proposal_data.get('support_type', 'N/A')),
            ('지원 금액', proposal_data.get('support_amount', 'N/A'))
        ]

        for i, (label, value) in enumerate(project_data):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = str(value)

        doc.add_paragraph()

    def _add_section_content(self, doc: Document, content: str):
        """섹션 내용 추가"""

        # 문단 분리
        paragraphs = content.split('\n\n')

        for para_text in paragraphs:
            if not para_text.strip():
                continue

            # 하위 제목 처리 (### 또는 ##로 시작)
            if para_text.strip().startswith('###'):
                text = para_text.strip().replace('###', '').strip()
                heading = doc.add_heading(text, level=3)
                heading.paragraph_format.space_before = Pt(6)
                heading.paragraph_format.space_after = Pt(3)

            elif para_text.strip().startswith('##'):
                text = para_text.strip().replace('##', '').strip()
                heading = doc.add_heading(text, level=2)
                heading.paragraph_format.space_before = Pt(9)
                heading.paragraph_format.space_after = Pt(6)

            # 리스트 항목 처리
            elif para_text.strip().startswith('- '):
                lines = para_text.split('\n')
                for line in lines:
                    if line.strip().startswith('- '):
                        text = line.strip()[2:]  # "- " 제거
                        para = doc.add_paragraph(text, style='List Bullet')
                        para.paragraph_format.space_after = Pt(3)

            # 번호 리스트 처리
            elif para_text.strip()[0].isdigit() and para_text.strip()[1:3] in ['. ', ') ']:
                lines = para_text.split('\n')
                for line in lines:
                    if line.strip() and line.strip()[0].isdigit():
                        # "1. " 또는 "1) " 형태에서 텍스트 추출
                        text = line.strip().split('. ', 1)[1] if '. ' in line else line.strip().split(') ', 1)[1]
                        para = doc.add_paragraph(text, style='List Number')
                        para.paragraph_format.space_after = Pt(3)

            # 일반 문단
            else:
                para = doc.add_paragraph(para_text.strip())
                para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                para.paragraph_format.space_after = Pt(6)
                para.paragraph_format.line_spacing = 1.5

    def generate_from_sections(
        self,
        title: str,
        sections: Dict[str, str],
        metadata: Optional[Dict] = None,
        output_path: str = None
    ) -> str:
        """
        섹션 딕셔너리로부터 DOCX 생성 (간단한 버전)

        Args:
            title: 문서 제목
            sections: 섹션 딕셔너리 {제목: 내용}
            metadata: 메타데이터 (선택)
            output_path: 출력 경로

        Returns:
            생성된 파일 경로
        """

        if not output_path:
            output_path = f"{title.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.docx"

        doc = Document()
        self._set_document_font(doc)

        # 제목
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # 메타데이터 (있는 경우)
        if metadata:
            for key, value in metadata.items():
                para = doc.add_paragraph()
                para.add_run(f"{key}: ").bold = True
                para.add_run(str(value))

            doc.add_paragraph()

        # 섹션별 내용
        for section_name, section_content in sections.items():
            doc.add_heading(section_name, level=1)
            self._add_section_content(doc, section_content)
            doc.add_paragraph()

        doc.save(output_path)
        return output_path
